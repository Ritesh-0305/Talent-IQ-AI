import io
import os
from typing import Tuple, Dict, Any
import fitz  # PyMuPDF
from docx import Document

class ExtractionError(Exception):
    pass

class LowTextWarning(Exception):
    pass

def extract_text_from_pdf(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    """
    Extracts text and page metadata from PDF bytes using PyMuPDF (fitz).
    Detects scanned/image-only PDFs and empty files.
    """
    if not file_bytes:
        raise ExtractionError("Cannot extract text from empty document (0 bytes).")
    
    # Check PDF magic bytes
    if not file_bytes.startswith(b"%PDF"):
        raise ExtractionError("Invalid PDF format: file header does not match %PDF.")

    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception as e:
        raise ExtractionError(f"Malformed or corrupted PDF file: {str(e)}")

    full_text = []
    page_count = len(doc)
    
    if page_count == 0:
        raise ExtractionError("PDF document has 0 pages.")

    for page_num in range(page_count):
        page = doc.load_page(page_num)
        page_text = page.get_text("text")
        if page_text:
            full_text.append(page_text.strip())

    combined_text = "\n\n".join(full_text).strip()
    
    # Check extractable text length
    # Note: OCR is explicitly not pretended unless an OCR engine exists.
    is_low_text = len(combined_text) < 30
    
    metadata = {
        "page_count": page_count,
        "char_count": len(combined_text),
        "word_count": len(combined_text.split()),
        "low_text_warning": is_low_text,
        "warning_message": "Scanned document or minimal extractable text detected. Optical Character Recognition (OCR) is not supported for raster images." if is_low_text else None
    }
    
    return combined_text, metadata


def extract_text_from_docx(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    """
    Extracts text from DOCX bytes using python-docx.
    Validates ZIP/PK format and paragraph/table structure.
    """
    if not file_bytes:
        raise ExtractionError("Cannot extract text from empty document (0 bytes).")
    
    # Check PK zip header
    if not file_bytes.startswith(b"PK\x03\x04"):
        raise ExtractionError("Invalid DOCX format: missing standard ZIP/DOCX header signature.")

    try:
        doc = Document(io.BytesIO(file_bytes))
    except Exception as e:
        raise ExtractionError(f"Malformed or corrupted DOCX file: {str(e)}")

    text_parts = []
    for paragraph in doc.paragraphs:
        if paragraph.text and paragraph.text.strip():
            text_parts.append(paragraph.text.strip())

    # Extract text from tables if any
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                text_parts.append(" | ".join(row_text))

    combined_text = "\n".join(text_parts).strip()
    is_low_text = len(combined_text) < 30

    metadata = {
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "char_count": len(combined_text),
        "word_count": len(combined_text.split()),
        "low_text_warning": is_low_text,
        "warning_message": "Document contains little or no extractable text." if is_low_text else None
    }

    return combined_text, metadata


def extract_resume_text(filename: str, file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    ext = os.path.splitext(filename.lower())[1]
    if ext == ".pdf":
        return extract_text_from_pdf(file_bytes)
    elif ext in [".docx", ".doc"]:
        if ext == ".doc":
            raise ExtractionError("Legacy binary .doc format is not supported. Please convert your resume to .docx or .pdf.")
        return extract_text_from_docx(file_bytes)
    else:
        raise ExtractionError(f"Unsupported file format '{ext}'. Only .pdf and .docx documents are accepted.")
